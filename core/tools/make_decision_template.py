"""Sinh mẫu Quyết định cử nhân sự (DOIT_QD_CuNhanSu_template.docx) với các thẻ poi-tl {{...}}."""
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

out = sys.argv[1]
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
st = d.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.15

def para(container, text="", bold=False, italic=False, align=None, size=None, before=0, after=0, indent=None):
    p = container.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    if indent is not None: p.paragraph_format.first_line_indent = Cm(indent)
    if text:
        r = p.add_run(text); r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
    return p

def runs(p, parts):
    """parts: list of (text, bold, italic)"""
    for t, b, i in parts:
        r = p.add_run(t); r.bold, r.italic = b, i
    return p

# Quốc hiệu / tên cơ quan
t = d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
w = [Cm(6.3), Cm(9.7)]
for i, c in enumerate(t.rows[0].cells): c.width = w[i]
for i, g in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))): g.set(qn("w:w"), str(int(w[i].twips)))
L, R = t.rows[0].cells
L.paragraphs[0].text = ""
runs(L.paragraphs[0], [("CÔNG TY TNHH CÔNG NGHỆ", True, False)]).alignment = A.CENTER
para(L, "VÀ TRUYỀN THÔNG DO IT", bold=True, align=A.CENTER)
para(L, "———————", align=A.CENTER)
runs(para(L, align=A.CENTER), [("Số: ", False, False), ("{{soQuyetDinh}}", False, False)])
runs(R.paragraphs[0], [("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", True, False)]).alignment = A.CENTER
para(R, "Độc lập – Tự do – Hạnh phúc", bold=True, align=A.CENTER)
para(R, "———————", align=A.CENTER)
runs(para(R, align=A.CENTER, before=4), [("Hà Nội, ngày {{ngay}} tháng {{thang}} năm {{nam}}", False, True)])

para(d, "QUYẾT ĐỊNH", bold=True, align=A.CENTER, size=14, before=18)
runs(para(d, align=A.CENTER), [("Về việc {{trichYeu}}", True, False)])
para(d, "———————", align=A.CENTER, after=6)
para(d, "GIÁM ĐỐC CÔNG TY TNHH CÔNG NGHỆ VÀ TRUYỀN THÔNG DO IT", bold=True, align=A.CENTER, after=6)

for line in ["Căn cứ Bộ luật Lao động số 45/2019/QH14 ngày 20/11/2019;",
             "Căn cứ Điều lệ tổ chức và hoạt động của Công ty;",
             "Căn cứ Hợp đồng lao động số {{soHopDong}};",
             "Căn cứ nhu cầu công việc của Công ty,"]:
    para(d, line, italic=True, indent=1, align=A.JUSTIFY)

para(d, "QUYẾT ĐỊNH:", bold=True, align=A.CENTER, before=8, after=6)

p = para(d, align=A.JUSTIFY, indent=1, after=3)
runs(p, [("Điều 1. ", True, False), ("Cử Ông/Bà ", False, False), ("{{hoTen}}", True, False),
         (", mã nhân viên {{maNhanVien}}, chức danh {{viTri}}, {{boPhan}}, {{hinhThuc}} ", False, False),
         ("{{donVi}}", True, False), (".", False, False)])
para(d, "– Địa chỉ: {{diaChi}}.", indent=1, align=A.JUSTIFY)
para(d, "– Thời gian: từ ngày {{tuNgay}} đến ngày {{denNgay}}.", indent=1, align=A.JUSTIFY)
para(d, "– Nội dung công việc: {{noiDung}}.", indent=1, align=A.JUSTIFY, after=3)
p = para(d, align=A.JUSTIFY, indent=1, after=3)
runs(p, [("Điều 2. ", True, False), ("Chế độ: {{cheDo}}.", False, False)])
p = para(d, align=A.JUSTIFY, indent=1, after=3)
runs(p, [("Điều 3. ", True, False),
         ("Quyết định có hiệu lực kể từ ngày ký. Bộ phận Hành chính – Nhân sự, Kế toán và Ông/Bà có tên tại Điều 1 "
          "chịu trách nhiệm thi hành Quyết định này.", False, False)])

t2 = d.add_table(rows=1, cols=2); t2.autofit = False
for i, c in enumerate(t2.rows[0].cells): c.width = [Cm(7.5), Cm(8.5)][i]
L, R = t2.rows[0].cells
r = L.paragraphs[0].add_run("Nơi nhận:"); r.bold, r.italic = True, True; r.font.size = Pt(12)
for x in ["– Như Điều 3;", "– Lưu: HCNS."]:
    q = para(L, x); q.runs[0].font.size = Pt(11)
R.paragraphs[0].add_run("{{chucVuKy}}").bold = True
R.paragraphs[0].alignment = A.CENTER
para(R, "(Ký, ghi rõ họ tên, đóng dấu)", italic=True, align=A.CENTER)
for _ in range(4): para(R)
para(R, "{{nguoiKy}}", bold=True, align=A.CENTER)
for ci, c in enumerate(t.rows[0].cells):
    for q in c.paragraphs:
        for r in q.runs: r.font.size = Pt(11) if ci == 0 else Pt(12)
# đẩy bảng ký xuống một chút
t2._tbl.addprevious(para(d, before=10)._p)
d.save(out)
print("saved", out)
