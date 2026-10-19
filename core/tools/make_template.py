"""Chuyển mẫu HĐLĐ của DOIT thành template poi-tl ({{truong}}), giữ nguyên định dạng."""
import copy
import sys

import docx
import docx.enum.text

src, dst = sys.argv[1], sys.argv[2]
d = docx.Document(src)
P = d.paragraphs


def set_runs(par, texts):
    """Gán lại text cho các run hiện có; run thừa thì xóa, thiếu thì nhân bản run cuối."""
    runs = par.runs
    while len(runs) < len(texts):
        new = copy.deepcopy(runs[-1]._r)
        runs[-1]._r.addnext(new)
        runs = par.runs
    for r, t in zip(runs, texts):
        r.text = t
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)


def fill_empty_cell(cell, text, style_from):
    """Thêm run vào ô trống, copy định dạng chữ từ ô nhãn bên cạnh."""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    src_r = style_from.paragraphs[0].runs[0]._r
    if src_r.rPr is not None:
        r._r.insert(0, copy.deepcopy(src_r.rPr))


# --- Số hợp đồng, người lao động
set_runs(P[5], ["{{soHopDong}}"])
set_runs(P[12], ["NGƯỜI LAO ĐỘNG:", " {{hoTen}}"])

# --- Bảng thông tin người lao động
t = d.tables[1]
lab = t.rows[0].cells[0]
fill_empty_cell(t.rows[0].cells[2], "{{noiCuTru}}", lab)
fill_empty_cell(t.rows[1].cells[2], "{{dienThoai}}", lab)
set_runs(t.rows[1].cells[3].paragraphs[0], ["Email:", " {{email}}"])
fill_empty_cell(t.rows[2].cells[2], "{{ngaySinh}}", lab)
set_runs(t.rows[2].cells[3].paragraphs[0], ["Giới tính: {{gioiTinh}}    Trình độ học vấn: {{trinhDo}}"])
set_runs(t.rows[3].cells[0].paragraphs[0], ["Số CCCD/Căn cước"])
fill_empty_cell(t.rows[3].cells[2], "{{soCanCuoc}}", lab)
set_runs(t.rows[3].cells[3].paragraphs[0], ["Cấp ngày: {{ngayCap}}", "   ", "Tại: ", "{{noiCap}}"])
fill_empty_cell(t.rows[4].cells[2], "{{soTaiKhoan}}", lab)
set_runs(t.rows[4].cells[3].paragraphs[0], ["Tại:", " {{nganHang}}"])

# --- Điều 1
set_runs(P[16], ["Loại hợp đồng: ", "{{loaiHopDong}}", "."])
set_runs(P[17], ["Thời gian hợp đồng: {{thoiHan}} bắt đầu từ {{tuNgay}} đến {{denNgay}}."])
set_runs(P[20], ["Thuộc ", "bộ phận", ": {{boPhan}}"])
set_runs(P[21], ["Vị trí: ", "{{viTri}}."])

# --- Điều 3A: bổ sung mức lương / phụ cấp / nâng lương (Điều 21 BLLĐ 2019)
# Dòng "Mức thu nhập (đã bao gồm phụ cấp)" gốc không có số tiền -> thay bằng 3 dòng cùng kiểu
# với dòng "Hình thức trả lương" (nhãn đậm + giá trị thường).
tmpl = P[35]  # "Hình thức trả lương: ..." làm khuôn
anchor = P[34]
lines = [
    ("Mức lương theo công việc", ": {{mucLuong}} đồng/tháng (Bằng chữ: {{mucLuongBangChu}})."),
    ("Phụ cấp lương và các khoản bổ sung khác", ": {{phuCap}}."),
    ("Chế độ nâng bậc, nâng lương", ": {{cheDoNangLuong}}."),
]
prev = anchor._p
for label, value in lines:
    new_p = copy.deepcopy(tmpl._p)
    prev.addnext(new_p)
    prev = new_p
    par = docx.text.paragraph.Paragraph(new_p, anchor._parent)
    set_runs(par, [label, value])
anchor._p.getparent().remove(anchor._p)

# --- Ô ký tên
sig = d.tables[2].rows[0].cells
for cell, name in ((sig[0], "{{nguoiDaiDien}}"), (sig[1], "{{hoTen}}")):
    for _ in range(4):
        cell.add_paragraph()
    p = cell.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name)
    r.bold = True

d.save(dst)
print("saved", dst)
